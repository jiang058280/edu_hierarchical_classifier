"""资料顺序与表格空列的回归测试，不访问外部服务。"""

from io import BytesIO

from docx import Document
import pytest

from edu_core.rag.chunking import build_parent_child_chunks
from edu_core.rag.loaders import load_document_bytes


def test_docx_preserves_interleaved_paragraphs_and_tables():
    document = Document()
    document.add_paragraph("第一节条件")
    table = document.add_table(rows=1, cols=3)
    table.cell(0, 0).text = "项目"
    table.cell(0, 2).text = "结论"
    document.add_paragraph("第二节说明")
    second = document.add_table(rows=1, cols=1)
    second.cell(0, 0).text = "第二张表"
    document.add_paragraph("最后总结")
    stream = BytesIO()
    document.save(stream)

    text = load_document_bytes("交错资料.docx", stream.getvalue()).text

    assert text.split("\n\n") == [
        "第一节条件", "项目 | | 结论", "第二节说明", "第二张表", "最后总结",
    ]


def test_docx_skips_empty_table_rows():
    document = Document()
    document.add_table(rows=1, cols=2)
    document.add_paragraph("有效正文")
    stream = BytesIO()
    document.save(stream)
    assert load_document_bytes("空表.docx", stream.getvalue()).text == "有效正文"


def test_chunks_do_not_assign_previous_text_to_next_chapter():
    chunks = build_parent_child_chunks(
        "# 第一章\n\n旧章内容\n\n# 第二章\n\n新章内容",
        parent_chars=100, child_chars=50, overlap_chars=10,
    )
    parents = [item for item in chunks if item.chunk_kind == "parent"]
    assert [(item.chapter, item.content) for item in parents] == [
        ("第一章", "# 第一章\n\n旧章内容"), ("第二章", "# 第二章\n\n新章内容"),
    ]
    assert all(item.chapter == "第一章" for item in chunks if "旧章内容" in item.content)


def test_long_paragraph_does_not_overtake_short_buffer():
    long_text = "长段内容。" * 35
    chunks = build_parent_child_chunks(
        "前置条件\n\n" + long_text + "\n\n最后说明",
        parent_chars=100, child_chars=50,
    )
    parents = [item.content for item in chunks if item.chunk_kind == "parent"]
    assert parents[0] == "前置条件"
    assert parents[-1] == "最后说明"
    assert "".join(parents[1:-1]) == long_text


def test_overlap_respects_parent_limit_and_preserves_new_paragraph():
    chunks = build_parent_child_chunks(
        "甲" * 90 + "\n\n" + "乙" * 95 + "\n\n" + "丙" * 20,
        parent_chars=100, child_chars=50, overlap_chars=30,
    )
    parents = [item.content for item in chunks if item.chunk_kind == "parent"]
    assert all(len(item) <= 100 for item in parents)
    assert any("乙" * 95 in item for item in parents)
    assert parents[-1].endswith("丙" * 20)


@pytest.mark.parametrize("overlap", [-1, 100, 101])
def test_invalid_overlap_is_rejected(overlap):
    with pytest.raises(ValueError, match="重叠长度"):
        build_parent_child_chunks("正文", parent_chars=100, child_chars=50, overlap_chars=overlap)
