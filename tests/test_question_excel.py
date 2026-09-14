"""F1 Excel 题库批量通道测试。"""

from io import BytesIO

from openpyxl import load_workbook

from edu_core.application.question_excel import (
    HEADERS,
    build_export,
    build_template,
    parse_workbook,
)


def _filled_template() -> bytes:
    workbook = load_workbook(BytesIO(build_template()))
    sheet = workbook["题库导入"]
    values = [
        "关于一元二次方程根的判断，下列说法正确的是？", "数学", "单选题", "一元二次方程",
        "初中", "初三", 2, "B", "判别式大于零时有两个不等实根。",
        "A. 没有实根|B. 有两个不等实根|C. 只有一个实根|D. 无法判断", "published",
    ]
    for index, value in enumerate(values, start=1):
        sheet.cell(2, index).value = value
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_template_has_expected_headers_and_validation():
    workbook = load_workbook(BytesIO(build_template()))
    sheet = workbook["题库导入"]
    assert tuple(cell.value for cell in sheet[1]) == HEADERS
    assert len(sheet.data_validations.dataValidation) >= 5
    assert sheet.freeze_panes == "A2"


def test_parse_valid_question_and_structured_options():
    rows, errors = parse_workbook(_filled_template())
    assert not errors
    assert len(rows) == 1
    assert rows[0]["options"][1] == {"key": "B", "text": "有两个不等实根"}
    assert rows[0]["difficulty"] == 2


def test_formula_cell_is_rejected():
    workbook = load_workbook(BytesIO(_filled_template()))
    sheet = workbook["题库导入"]
    sheet["A2"] = '=HYPERLINK("https://example.invalid","题目")'
    output = BytesIO()
    workbook.save(output)

    rows, errors = parse_workbook(output.getvalue())

    assert not rows
    assert errors == [{"row": 2, "field": "text", "message": "不允许使用公式单元格"}]


def test_export_can_be_imported_again_and_text_is_not_formula():
    payload = build_export([{
        "content": "=这是一道以等号开头但长度足够的普通题干", "subject": "数学",
        "question_type": "填空题", "knowledge_point": "代数", "grade_band": "高中",
        "grade": "高一", "difficulty": 3, "answer": "1", "analysis": "解析",
        "options": None, "status": "draft",
    }])
    workbook = load_workbook(BytesIO(payload), data_only=False)
    assert workbook.active["A2"].data_type == "s"
    rows, errors = parse_workbook(payload)
    assert not errors
    assert rows[0]["content"].startswith("=")
