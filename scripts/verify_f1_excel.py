"""F1 Excel 题库通道本机验收：模板、导出、100 行异步回导与幂等。"""

from __future__ import annotations

import json
import sys
import time
from datetime import datetime
from io import BytesIO
from pathlib import Path

import httpx
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from edu_core.application.question_excel import HEADERS
from edu_core.security.auth import create_access_token
from edu_core.storage.stores import StoreBundle


def _trim_to_100(payload: bytes) -> bytes:
    workbook = load_workbook(BytesIO(payload))
    sheet = workbook.active
    if sheet.max_row < 101:
        raise RuntimeError("题库不足 100 行，无法执行容量验收")
    if sheet.max_row > 101:
        sheet.delete_rows(102, sheet.max_row - 101)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def main() -> int:
    stores = StoreBundle()
    user = stores.users.get_by_username("lgq")
    if not user or user["role"] not in {"teacher", "admin"}:
        raise RuntimeError("教师账号 lgq 不存在")
    token = create_access_token(user["username"], user["role"], int(user["id"]))
    headers = {"Authorization": f"Bearer {token}"}
    started = time.perf_counter()
    with httpx.Client(base_url="http://127.0.0.1:7860", headers=headers, timeout=20) as client:
        template = client.get("/api/v1/teacher/questions/template")
        template.raise_for_status()
        template_book = load_workbook(BytesIO(template.content), read_only=True)
        template_headers = tuple(cell.value for cell in next(template_book.active.iter_rows(max_row=1)))
        if template_headers != HEADERS:
            raise RuntimeError("模板表头与导入契约不一致")

        # 历史题库由本次结构化资料导入，数量超过 100 且字段完整，适合作为可回导夹具。
        exported = client.get(
            "/api/v1/teacher/questions/export", params={"subject": "历史", "status": "published"})
        exported.raise_for_status()
        workbook_100 = _trim_to_100(exported.content)

        import_started = time.perf_counter()
        submitted = client.post(
            "/api/v1/teacher/questions/import",
            files={"file": ("回导验收.xlsx", workbook_100,
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        submitted.raise_for_status()
        job_id = submitted.json()["job_id"]
        deadline = time.monotonic() + 10
        while True:
            job_response = client.get(f"/api/v1/teacher/questions/import/{job_id}")
            job_response.raise_for_status()
            job = job_response.json()
            if job["status"] in {"SUCCEEDED", "FAILED"}:
                break
            if time.monotonic() >= deadline:
                raise RuntimeError("100 行导入任务 10 秒内未完成")
            time.sleep(0.1)
        import_elapsed = time.perf_counter() - import_started
        if job["status"] != "SUCCEEDED":
            raise RuntimeError(f"导入任务失败：{job.get('errors')}")
        if job["total"] != 100 or job["inserted"] != 0 or job["skipped"] != 100 or job["failed"] != 0:
            raise RuntimeError(f"幂等回导结果不符合预期：{job}")
        if import_elapsed >= 5:
            raise RuntimeError(f"100 行导入耗时 {import_elapsed:.3f}s，不满足 <5s")

    report = {
        "verified_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "status": "passed",
        "template_headers": list(template_headers),
        "export_bytes": len(exported.content),
        "import_rows": 100,
        "inserted": job["inserted"],
        "skipped_existing": job["skipped"],
        "failed": job["failed"],
        "import_elapsed_seconds": round(import_elapsed, 3),
        "total_elapsed_seconds": round(time.perf_counter() - started, 3),
    }
    output = ROOT / "reports" / "verification" / "f1_excel_latest.json"
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
