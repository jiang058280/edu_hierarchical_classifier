"""教师端路由组（平台计划 M0/M1/M2，require_teacher 门禁）。

M0：班级管理（建班/名单/邀请码）。
M1：题库完整 CRUD 与筛选、AI 预标注（单/批）、组卷生成/保存/Word 导出。
M2：作业发布、详情、提交进度与教师批改。
"""

from __future__ import annotations

import secrets
from io import BytesIO
from typing import Any
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import Response

from edu_core.application.entry_annotation import (
    extract_embedded_metadata,
    normalize_knowledge_point,
    suggest_difficulty,
)
from edu_core.security.auth import require_teacher
from edu_core.storage.stores import StoreBundle

router = APIRouter()


def _stores() -> StoreBundle:
    return StoreBundle()


# ---------------------------------------------------------------------------
# 教学知识库（R1.5）：教师资料管理与版本发布
# ---------------------------------------------------------------------------

@router.get("/teacher/rag/versions")
def list_rag_versions(user: dict[str, Any] = Depends(require_teacher)) -> dict:
    stores = _stores()
    return {"items": stores.rag.list_versions(created_by=int(user["id"])),
            "active": stores.rag.get_active_version()}


@router.post("/teacher/rag/versions")
def create_rag_version(payload: dict, request: Request,
                       user: dict[str, Any] = Depends(require_teacher)) -> dict:
    version = (payload.get("version") or "").strip()
    description = (payload.get("description") or "").strip()
    try:
        version_id = _stores().rag.create_version(version, created_by=int(user["id"]), description=description)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    _stores().audit.insert(action="create_rag_version", user_id=int(user["id"]), username=user.get("username"),
                           resource=f"rag/versions/{version_id}", detail={"version": version},
                           client_ip=request.client.host if request.client else None)
    return {"status": "ok", "id": version_id}


@router.delete("/teacher/rag/versions/{version_id}")
def delete_rag_version(version_id: int, request: Request,
                       user: dict[str, Any] = Depends(require_teacher)) -> dict:
    """删除教师自己的非激活版本及关联资料；当前激活版本受保护。"""
    from edu_core.config.settings import get_settings
    from edu_core.rag.indexing.milvus_index import MilvusRagDocumentIndex

    stores, settings = _stores(), get_settings()
    version = stores.rag.get_owned_version(version_id, int(user["id"]))
    if not version:
        raise HTTPException(status_code=404, detail="知识库版本不存在")
    documents = stores.rag.list_documents(created_by=int(user["id"]), kb_version_id=version_id)
    chunk_ids = [int(chunk["id"]) for document in documents
                 for chunk in stores.rag.list_document_chunks(int(document["id"]))]
    MilvusRagDocumentIndex(settings).delete_chunks(chunk_ids)
    try:
        cleanup = stores.rag.delete_version(version_id, created_by=int(user["id"]))
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    upload_root = settings.abs_path(settings.rag_upload_dir).resolve()
    for storage_key in cleanup["storage_keys"]:
        source_path = settings.abs_path(str(storage_key)).resolve()
        if upload_root in source_path.parents:
            try:
                source_path.unlink(missing_ok=True)
            except OSError:
                pass
    stores.audit.insert(action="delete_rag_version", user_id=int(user["id"]), username=user.get("username"),
                        resource=f"rag/versions/{version_id}",
                        detail={"version": version["version"], "documents": cleanup["document_count"], "permanent": True},
                        client_ip=request.client.host if request.client else None)
    return {"status": "ok", "deleted_version_id": version_id,
            "deleted_documents": cleanup["document_count"]}


@router.get("/teacher/rag/documents")
def list_rag_documents(kb_version_id: int | None = None,
                       user: dict[str, Any] = Depends(require_teacher)) -> dict:
    return {"items": _stores().rag.list_documents(created_by=int(user["id"]), kb_version_id=kb_version_id)}


@router.get("/teacher/rag/question-knowledge")
def question_knowledge_status(_: dict[str, Any] = Depends(require_teacher)) -> dict:
    return _stores().rag.question_knowledge_status()


@router.post("/teacher/rag/question-knowledge/sync")
def sync_question_knowledge(request: Request,
                            user: dict[str, Any] = Depends(require_teacher)) -> dict:
    result = _stores().rag.sync_published_questions()
    _stores().audit.insert(action="sync_rag_question_knowledge", user_id=int(user["id"]),
                           username=user.get("username"), resource="rag/question-knowledge",
                           detail=result | {"external_model_called": False},
                           client_ip=request.client.host if request.client else None)
    return {"status": "ok", **result, "external_model_called": False}


@router.get("/teacher/rag/question-knowledge/search")
def search_question_knowledge(query: str, subject: str = "", grade_band: str = "", grade: str = "",
                              _: dict[str, Any] = Depends(require_teacher)) -> dict:
    """教师验证本地题库知识源；不构造 Embedding，不调用外部模型。"""
    from edu_core.config.settings import get_settings
    from edu_core.rag.retrieval import RagRetrievalService, RetrievalFilters

    query = query.strip()
    if not query:
        raise HTTPException(status_code=400, detail="请输入要检索的问题")
    hits = RagRetrievalService(_stores().rag, get_settings()).local_question_candidates(
        query, filters=RetrievalFilters(subject=subject.strip() or None, grade_band=grade_band.strip() or None,
                                        grade=grade.strip() or None))
    return {"items": hits, "external_model_called": False}


@router.get("/teacher/rag/documents/{document_id}")
def preview_rag_document(document_id: int, user: dict[str, Any] = Depends(require_teacher)) -> dict:
    document = _stores().rag.get_owned_document(document_id, int(user["id"]))
    if not document:
        raise HTTPException(status_code=404, detail="资料不存在")
    chunks = _stores().rag.list_document_chunks(document_id)
    return {"document": document,
            "ocr_review": _stores().rag.get_ocr_review(document_id, created_by=int(user["id"])),
            "chunk_summary": {
        "parent_count": sum(chunk["chunk_kind"] == "parent" for chunk in chunks),
        "child_count": sum(chunk["chunk_kind"] == "child" for chunk in chunks),
        "sample": [chunk["content"][:240] for chunk in chunks if chunk["chunk_kind"] == "child"][:3],
    }}


@router.post("/teacher/rag/documents/upload")
async def upload_rag_document(request: Request, file: UploadFile = File(...), kb_version_id: int = Form(...),
                              subject: str = Form(""), grade_band: str = Form(""), grade: str = Form(""),
                              knowledge_node_id: int | None = Form(None),
                              user: dict[str, Any] = Depends(require_teacher)) -> dict:
    from edu_core.config.settings import get_settings
    from edu_core.rag.ingestion import RagIngestionService

    source_name = (file.filename or "").strip()
    if not source_name:
        raise HTTPException(status_code=400, detail="请选择资料文件")
    settings = get_settings()
    content = await file.read(int(settings.rag_max_upload_bytes) + 1)
    if len(content) > int(settings.rag_max_upload_bytes):
        raise HTTPException(status_code=400, detail="资料超过上传大小限制")
    stores = _stores()
    if not stores.rag.get_owned_version(kb_version_id, int(user["id"])):
        raise HTTPException(status_code=404, detail="知识库版本不存在")
    try:
        result = RagIngestionService(stores.rag, settings).ingest(
            kb_version_id=kb_version_id, created_by=int(user["id"]), source_name=source_name, content=content,
            subject=subject.strip() or None, grade_band=grade_band.strip() or None, grade=grade.strip() or None,
            knowledge_node_id=knowledge_node_id)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:  # noqa: BLE001 - 任务本身已记录为 FAILED，接口仅返回受控信息
        raise HTTPException(status_code=503, detail=f"资料入库失败：{str(exc)[:160]}") from exc
    stores.audit.insert(action="upload_rag_document", user_id=int(user["id"]), username=user.get("username"),
                        resource=f"rag/documents/{result.document_id}",
                        detail={"source_name": source_name, "reused": result.reused},
                        client_ip=request.client.host if request.client else None)
    return {"status": "ok", "document_id": result.document_id, "job_id": result.job_id,
            "chunk_count": result.chunk_count, "reused": result.reused}


@router.post("/teacher/rag/documents/{document_id}/rebuild")
def rebuild_rag_document(document_id: int, request: Request,
                         user: dict[str, Any] = Depends(require_teacher)) -> dict:
    """以已保存的原始资料重建为新记录，旧已发布资料不被覆盖。"""
    from edu_core.config.settings import get_settings
    from edu_core.rag.ingestion import RagIngestionService

    stores, settings = _stores(), get_settings()
    document = stores.rag.get_owned_document(document_id, int(user["id"]))
    if not document or not document.get("storage_key"):
        raise HTTPException(status_code=404, detail="可重建的资料不存在")
    upload_root = settings.abs_path(settings.rag_upload_dir).resolve()
    source_path = settings.abs_path(document["storage_key"]).resolve()
    if upload_root not in source_path.parents or not source_path.is_file():
        raise HTTPException(status_code=404, detail="原始资料文件不存在")
    try:
        result = RagIngestionService(stores.rag, settings).ingest(
            kb_version_id=int(document["kb_version_id"]), created_by=int(user["id"]),
            source_name=document["source_name"], content=source_path.read_bytes(), subject=document.get("subject"),
            grade_band=document.get("grade_band"), grade=document.get("grade"),
            knowledge_node_id=document.get("knowledge_node_id"), allowed_roles=document.get("allowed_roles"),
            allow_reuse=False)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=503, detail=f"资料重建失败：{str(exc)[:160]}") from exc
    stores.audit.insert(action="rebuild_rag_document", user_id=int(user["id"]), username=user.get("username"),
                        resource=f"rag/documents/{result.document_id}", detail={"source_document_id": document_id},
                        client_ip=request.client.host if request.client else None)
    return {"status": "ok", "document_id": result.document_id, "job_id": result.job_id,
            "chunk_count": result.chunk_count}


@router.get("/teacher/rag/documents/{document_id}/source")
def download_rag_source(document_id: int, user: dict[str, Any] = Depends(require_teacher)):
    from fastapi.responses import FileResponse
    from edu_core.config.settings import get_settings
    settings = get_settings()
    doc = _stores().rag.get_owned_document(document_id, int(user["id"]))
    if not doc or not doc.get("storage_key"):
        raise HTTPException(status_code=404, detail="原件不存在")
    path = settings.abs_path(doc["storage_key"]).resolve()
    if settings.abs_path(settings.rag_upload_dir).resolve() not in path.parents or not path.is_file():
        raise HTTPException(status_code=404, detail="原件不存在")
    return FileResponse(path, filename=doc["source_name"], media_type="application/octet-stream")


@router.post("/teacher/rag/documents/{document_id}/correction")
def correct_rag_document(document_id: int, payload: dict, request: Request,
                         user: dict[str, Any] = Depends(require_teacher)) -> dict:
    """Create a reviewed-text draft with fresh embeddings; preserve the source record."""
    from pathlib import Path
    from edu_core.config.settings import get_settings
    from edu_core.rag.ingestion import RagIngestionService

    stores, settings = _stores(), get_settings()
    owner = int(user["id"])
    original = stores.rag.get_owned_document(document_id, owner)
    if not original:
        raise HTTPException(status_code=404, detail="资料不存在")
    corrected, note = payload.get("text"), payload.get("note")
    if not isinstance(corrected, str) or not corrected.strip() or len(corrected.encode("utf-8")) > min(settings.rag_max_upload_bytes, 500000):
        raise HTTPException(status_code=400, detail="校对文本不能为空且不得超过 500KB")
    if not isinstance(note, str) or not note.strip() or len(note) > 2000:
        raise HTTPException(status_code=400, detail="请填写 1–2000 字的校对说明")
    if original["status"] not in ("PROCESSED", "UNPUBLISHED", "PUBLISHED"):
        raise HTTPException(status_code=400, detail="仅已处理资料可以创建校对版")
    state = stores.rag.get_ocr_review(document_id, created_by=owner)
    if not state.get("token") or payload.get("token") != state["token"]:
        raise HTTPException(status_code=409, detail="资料内容已变化或尚未就绪，请重新预览")
    provenance = {"source_document_id": document_id, "source_content_hash": original["content_hash"],
                  "source_snapshot": state["token"], "corrected_by": owner, "note": note.strip()}
    try:
        result = RagIngestionService(stores.rag, settings).ingest(
            kb_version_id=int(original["kb_version_id"]), created_by=owner,
            source_name=Path(original["source_name"]).stem[:160] + "_在线校对.txt",
            content=corrected.strip().encode("utf-8"), subject=original.get("subject"),
            grade_band=original.get("grade_band"), grade=original.get("grade"),
            knowledge_node_id=original.get("knowledge_node_id"), allowed_roles=original.get("allowed_roles"),
            allow_reuse=False, correction_provenance=provenance)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=503, detail="校对版入库失败，请查看处理记录后重试") from exc
    stores.audit.insert(action="correct_rag_document", user_id=owner, username=user.get("username"),
                        resource=f"rag/documents/{result.document_id}", detail=provenance,
                        client_ip=request.client.host if request.client else None)
    return {"status": "ok", "document_id": result.document_id, "job_id": result.job_id,
            "chunk_count": result.chunk_count, "requires_review": True}


@router.post("/teacher/rag/documents/{document_id}/ocr-review")
def review_rag_ocr(document_id: int, payload: dict,
                   user: dict[str, Any] = Depends(require_teacher)) -> dict:
    """Persist an explicit human decision; never infer approval from OCR success."""
    if type(payload.get("approved")) is not bool:
        raise HTTPException(status_code=400, detail="approved 必须为布尔值")
    if payload["approved"] and payload.get("checked_source_title_table_math") is not True:
        raise HTTPException(status_code=400, detail="请确认已对照原件核对标题、表格及数学符号")
    try:
        result = _stores().rag.review_ocr_document(
            document_id, created_by=int(user["id"]), token=payload.get("token"),
            approved=payload["approved"], note=payload.get("note"))
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"status": "ok", "ocr_review": result}


@router.post("/teacher/rag/documents/{document_id}/publication")
def set_rag_document_publication(document_id: int, payload: dict, request: Request,
                                 user: dict[str, Any] = Depends(require_teacher)) -> dict:
    try:
        document = _stores().rag.set_document_publication(
            document_id, created_by=int(user["id"]), published=bool(payload.get("published")))
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    _stores().audit.insert(action="publish_rag_document" if payload.get("published") else "unpublish_rag_document",
                           user_id=int(user["id"]), username=user.get("username"),
                           resource=f"rag/documents/{document_id}", detail={"published": bool(payload.get("published"))},
                           client_ip=request.client.host if request.client else None)
    return {"status": "ok", "document": document}


@router.post("/teacher/rag/documents/{document_id}/archive")
def archive_rag_document(document_id: int, request: Request,
                         user: dict[str, Any] = Depends(require_teacher)) -> dict:
    try:
        document = _stores().rag.archive_document(document_id, created_by=int(user["id"]))
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    _stores().audit.insert(action="archive_rag_document", user_id=int(user["id"]), username=user.get("username"),
                           resource=f"rag/documents/{document_id}", detail={"archived": True},
                           client_ip=request.client.host if request.client else None)
    return {"status": "ok", "document": document}


@router.delete("/teacher/rag/documents/{document_id}")
def delete_rag_document(document_id: int, request: Request,
                        user: dict[str, Any] = Depends(require_teacher)) -> dict:
    """永久移除资料记录；先删向量，再删除 MySQL 关系数据与不再被引用的原文件。"""
    from edu_core.config.settings import get_settings
    from edu_core.rag.indexing.milvus_index import MilvusRagDocumentIndex

    stores, settings = _stores(), get_settings()
    document = stores.rag.get_owned_document(document_id, int(user["id"]))
    if not document:
        raise HTTPException(status_code=404, detail="资料不存在")
    chunk_ids = [int(chunk["id"]) for chunk in stores.rag.list_document_chunks(document_id)]
    try:
        MilvusRagDocumentIndex(settings).delete_chunks(chunk_ids)
        cleanup = stores.rag.delete_document(document_id, created_by=int(user["id"]))
    except ValueError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except Exception as exc:  # noqa: BLE001 - 向量清理失败时保留记录以便安全重试
        raise HTTPException(status_code=503, detail=f"资料删除失败：{str(exc)[:160]}") from exc
    if cleanup["delete_storage"]:
        upload_root = settings.abs_path(settings.rag_upload_dir).resolve()
        source_path = settings.abs_path(str(cleanup["storage_key"])).resolve()
        if upload_root in source_path.parents:
            try:
                source_path.unlink(missing_ok=True)
            except OSError:
                # 记录已删除，残留文件不能再被访问，后续可由运维按上传目录清理。
                pass
    stores.audit.insert(action="delete_rag_document", user_id=int(user["id"]), username=user.get("username"),
                        resource=f"rag/documents/{document_id}", detail={"permanent": True},
                        client_ip=request.client.host if request.client else None)
    return {"status": "ok", "deleted_document_id": document_id}


@router.post("/teacher/rag/versions/{version_id}/activate")
def activate_rag_version(version_id: int, request: Request,
                         user: dict[str, Any] = Depends(require_teacher)) -> dict:
    stores = _stores()
    try:
        report = stores.rag.build_quality_report(version_id, created_by=int(user["id"]))
        stores.rag.set_quality_report(version_id, report)
        version = stores.rag.activate_version(version_id)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    _stores().audit.insert(action="activate_rag_version", user_id=int(user["id"]), username=user.get("username"),
                           resource=f"rag/versions/{version_id}", detail={"report": report},
                           client_ip=request.client.host if request.client else None)
    return {"status": "ok", "version": version, "quality_report": report}


@router.post("/teacher/rag/retrieval/debug")
def rag_retrieval_debug(payload: dict, user: dict[str, Any] = Depends(require_teacher)) -> dict:
    """R2.1 检索调试：不调用 LLM，只返回经过 MySQL 授权过滤后的证据。"""
    from edu_core.config.settings import get_settings
    from edu_core.rag.retrieval import RagRetrievalService, RetrievalFilters

    filters = RetrievalFilters(
        subject=(payload.get("subject") or "").strip() or None,
        grade_band=(payload.get("grade_band") or "").strip() or None,
        grade=(payload.get("grade") or "").strip() or None,
        knowledge_node_id=int(payload["knowledge_node_id"]) if payload.get("knowledge_node_id") else None)
    try:
        return RagRetrievalService(_stores().rag, get_settings()).debug(
            payload.get("query") or "", role="teacher", filters=filters)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=503, detail=f"检索服务暂不可用：{str(exc)[:160]}") from exc


@router.post("/teacher/rag/answer/preview")
def rag_answer_preview(payload: dict, user: dict[str, Any] = Depends(require_teacher)) -> dict:
    """R2.2 教师预览：只基于已授权证据生成，返回服务端构造的引用。"""
    from edu_core.config.settings import get_settings
    from edu_core.rag.generation import RagAnswerService
    from edu_core.rag.retrieval import RagRetrievalService, RetrievalFilters

    filters = RetrievalFilters(
        subject=(payload.get("subject") or "").strip() or None,
        grade_band=(payload.get("grade_band") or "").strip() or None,
        grade=(payload.get("grade") or "").strip() or None,
        knowledge_node_id=int(payload["knowledge_node_id"]) if payload.get("knowledge_node_id") else None)
    settings, stores = get_settings(), _stores()
    try:
        result = RagAnswerService(RagRetrievalService(stores.rag, settings), settings).answer(
            payload.get("query") or "", role="teacher", filters=filters)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=503, detail=f"问答服务暂不可用：{str(exc)[:160]}") from exc
    return {"answer": result.answer, "citations": result.citations, "refused": result.refused,
            "reason": result.reason}


# ---------------------------------------------------------------------------
# 班级管理（M0）
# ---------------------------------------------------------------------------

@router.get("/teacher/classes")
def my_classes(user: dict[str, Any] = Depends(require_teacher)) -> dict:
    """我创建的班级列表（含学生数）。"""
    classes = _stores().classes.list_by_teacher(int(user["id"]))
    return {"classes": classes}


@router.post("/teacher/classes")
def create_class(payload: dict, user: dict[str, Any] = Depends(require_teacher)) -> dict:
    """创建班级：{name, grade_band, grade} → 生成 6 位邀请码。"""
    name = (payload.get("name") or "").strip()
    grade_band = (payload.get("grade_band") or "").strip()
    grade = (payload.get("grade") or "").strip()
    if not name or grade_band not in ("初中", "高中") or not grade:
        raise HTTPException(
            status_code=400, detail="name / grade(年级) 必填，grade_band 仅支持 初中/高中")
    invite_code = secrets.token_hex(3).upper()  # 6 位十六进制邀请码
    stores = _stores()
    try:
        class_id = stores.classes.create(
            name, grade_band, grade, created_by=int(user["id"]), invite_code=invite_code)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    stores.audit.insert(
        action="create_class", user_id=user.get("id") or None,
        username=user.get("username"), resource=f"classes/{class_id}",
        detail={"name": name, "grade_band": grade_band, "grade": grade})
    return {"status": "ok", "id": class_id, "invite_code": invite_code}


@router.get("/teacher/classes/{class_id}/students")
def class_students(class_id: int, user: dict[str, Any] = Depends(require_teacher)) -> dict:
    """班级学生名册（仅班级创建教师可看）。"""
    stores = _stores()
    cls = stores.classes.get(class_id)
    if not cls or cls["created_by"] != int(user["id"]):
        raise HTTPException(status_code=404, detail="班级不存在")
    return {"class": cls, "students": stores.users.list_by_class(class_id)}


@router.get("/teacher/analytics/class/{class_id}")
def class_learning_analytics(class_id: int, user: dict[str, Any] = Depends(require_teacher)) -> dict:
    """R3.6：只允许班级创建教师读取班级画像及重讲建议。"""
    from edu_core.application.factory import get_learning_service

    cls = _stores().classes.get(class_id)
    if not cls or int(cls["created_by"]) != int(user["id"]):
        raise HTTPException(status_code=404, detail="班级不存在")
    return {"class": cls, **get_learning_service().class_analytics(class_id)}


@router.get("/teacher/review-queue")
def review_queue(limit: int = 100, _: dict[str, Any] = Depends(require_teacher)) -> dict:
    items = _stores().classifications.review_queue(limit=max(1, min(limit, 200)))
    return {"items": items, "total": len(items)}


# ---------------------------------------------------------------------------
# 题库管理（M1）：完整题目 CRUD 与筛选
# ---------------------------------------------------------------------------

@router.get("/teacher/knowledge-nodes")
def knowledge_nodes(subject: str = "", level: int | None = None,
                    _: dict[str, Any] = Depends(require_teacher)) -> dict:
    """知识点树（录入页下拉与组卷筛选使用）。"""
    nodes = _stores().knowledge_nodes.list(subject=subject, level=level)
    return {"total": len(nodes), "nodes": nodes}


@router.get("/teacher/question-taxonomy")
def question_taxonomy(subject: str = "", grade_band: str = "", knowledge: str = "",
                      difficulty: int | None = None,
                      _: dict[str, Any] = Depends(require_teacher)) -> dict:
    """统一题型目录；选择学科时同时返回已发布题目的可用数量。"""
    from edu_core.application.question_taxonomy import taxonomy_payload

    counts = _stores().questions.question_type_counts(
        subject=subject, grade_band=grade_band, status="published",
        knowledge=knowledge, difficulty=difficulty)
    try:
        return taxonomy_payload(subject, counts)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.get("/teacher/questions")
def list_questions(subject: str = "", question_type: str = "", keyword: str = "",
                   grade_band: str = "", difficulty: int | None = None,
                   status: str = "", knowledge: str = "",
                   limit: int = 50, offset: int = 0,
                   _: dict[str, Any] = Depends(require_teacher)) -> dict:
    """题库筛选列表（学段/学科/题型/难度/状态/知识点/关键词）。"""
    items, total = _stores().questions.list(
        subject=subject, question_type=question_type, keyword=keyword,
        grade_band=grade_band, difficulty=difficulty, status=status,
        knowledge=knowledge, limit=min(limit, 200), offset=max(offset, 0))
    return {"total": total, "items": items}


def _xlsx_response(content: bytes, filename: str) -> Response:
    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@router.get("/teacher/questions/template")
def download_question_template(_: dict[str, Any] = Depends(require_teacher)) -> Response:
    """下载带字段说明和下拉校验的题库导入模板。"""
    from edu_core.application.question_excel import build_template

    try:
        return _xlsx_response(build_template(), "题库导入模板.xlsx")
    except RuntimeError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc


@router.post("/teacher/questions/import")
async def import_questions_xlsx(request: Request, file: UploadFile = File(...),
                                user: dict[str, Any] = Depends(require_teacher)) -> dict:
    """提交 Excel 异步导入任务；状态与错误行通过 job_id 查询。"""
    from edu_core.application.question_excel import get_question_import_jobs

    filename = (file.filename or "").strip()
    if not filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="仅支持 .xlsx 文件")
    max_bytes = 10 * 1024 * 1024
    payload = await file.read(max_bytes + 1)
    if not payload:
        raise HTTPException(status_code=400, detail="上传文件为空")
    if len(payload) > max_bytes:
        raise HTTPException(status_code=400, detail="Excel 文件不能超过 10 MB")
    jobs = get_question_import_jobs()
    job_id = jobs.submit(payload, user_id=int(user["id"]), filename=filename)
    _stores().audit.insert(
        action="import_questions_xlsx", user_id=int(user["id"]), username=user.get("username"),
        resource=f"teacher/questions/import/{job_id}", detail={"filename": filename},
        client_ip=request.client.host if request.client else None,
    )
    return {"status": "QUEUED", "job_id": job_id}


@router.get("/teacher/questions/import/{job_id}")
def question_import_status(job_id: str,
                           user: dict[str, Any] = Depends(require_teacher)) -> dict:
    from edu_core.application.question_excel import get_question_import_jobs

    job = get_question_import_jobs().get(job_id, int(user["id"]))
    if not job:
        raise HTTPException(status_code=404, detail="导入任务不存在或已随服务重启失效")
    return {key: value for key, value in job.items() if key != "user_id"}


@router.get("/teacher/questions/import/{job_id}/errors")
def download_question_import_errors(job_id: str,
                                    user: dict[str, Any] = Depends(require_teacher)) -> Response:
    from edu_core.application.question_excel import build_error_report, get_question_import_jobs

    job = get_question_import_jobs().get(job_id, int(user["id"]))
    if not job:
        raise HTTPException(status_code=404, detail="导入任务不存在或已随服务重启失效")
    if job["status"] not in {"SUCCEEDED", "FAILED"}:
        raise HTTPException(status_code=409, detail="导入任务尚未完成")
    return _xlsx_response(build_error_report(job.get("errors") or []), f"题库导入错误-{job_id[:8]}.xlsx")


@router.get("/teacher/questions/export")
def export_questions_xlsx(subject: str = "", question_type: str = "", keyword: str = "",
                          grade_band: str = "", difficulty: int | None = None,
                          status: str = "", knowledge: str = "",
                          _: dict[str, Any] = Depends(require_teacher)) -> Response:
    """按题库页面当前筛选条件导出完整题目信息，可直接再次导入。"""
    from edu_core.application.question_excel import build_export

    records = _stores().questions.export_rows(
        subject=subject, question_type=question_type, keyword=keyword,
        grade_band=grade_band, difficulty=difficulty, status=status, knowledge=knowledge,
    )
    return _xlsx_response(build_export(records), "题库导出.xlsx")


@router.post("/teacher/questions/import-docx")
async def import_questions_docx(file: UploadFile = File(...),
                                _: dict[str, Any] = Depends(require_teacher)) -> dict:
    """提取并切分 DOCX 题目，沿用 AI 预标注结果结构供教师逐题确认。"""
    from edu_core.application.question_docx import parse_docx_questions
    from edu_core.config.settings import get_settings

    filename = (file.filename or "").strip()
    if not filename.lower().endswith(".docx"):
        raise HTTPException(status_code=400, detail="仅支持 .docx 文件")
    max_bytes = 10 * 1024 * 1024
    payload = await file.read(max_bytes + 1)
    if not payload:
        raise HTTPException(status_code=400, detail="上传文件为空")
    if len(payload) > max_bytes:
        raise HTTPException(status_code=400, detail="DOCX 文件不能超过 10 MB")
    try:
        questions = parse_docx_questions(payload, max_questions=50)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    if not questions:
        raise HTTPException(status_code=400, detail="文档中未识别到有效题目")
    max_chars = int(get_settings().classify_max_chars)
    oversized = [index + 1 for index, value in enumerate(questions) if len(value) > max_chars]
    if oversized:
        raise HTTPException(status_code=400, detail=f"第 {oversized[0]} 题超过 {max_chars} 字符")
    service = _stores_service()
    items = [_preannotate(question, service) for question in questions]
    return {"items": items, "total": len(items), "filename": filename}


@router.get("/teacher/questions/{question_id}")
def get_question(question_id: int,
                 _: dict[str, Any] = Depends(require_teacher)) -> dict:
    """题目完整信息。"""
    question = _stores().questions.get(question_id)
    if not question:
        raise HTTPException(status_code=404, detail="题目不存在")
    return question


@router.delete("/teacher/questions/{question_id}")
def delete_teacher_question(question_id: int, request: Request,
                            user: dict[str, Any] = Depends(require_teacher)) -> dict:
    """删除题目（同步删除 Milvus 向量，动作写审计）。"""
    from edu_core.application.factory import get_classification_service

    service = get_classification_service()
    try:
        result = service.delete_question(question_id)
    except Exception as exc:
        from edu_core.application.service import ValidationError
        from edu_core.storage.stores import sqlalchemy_error_to_message
        if isinstance(exc, ValidationError):
            raise HTTPException(status_code=404, detail=str(exc))
        if isinstance(exc, ValueError):
            raise HTTPException(status_code=400, detail=str(exc)) from exc
        msg = sqlalchemy_error_to_message(exc)
        raise HTTPException(status_code=500, detail=msg or str(exc))
    _stores().audit.insert(
        action="delete_question", user_id=user.get("id") or None,
        username=user.get("username"), resource=f"teacher/questions/{question_id}",
        detail={"question_id": question_id}, client_ip=request.client.host if request.client else None)
    return result


def _full_fields(payload: dict, *, require_content: bool = True) -> dict:
    fields: dict[str, Any] = {}
    if require_content or payload.get("text"):
        text = (payload.get("text") or "").strip()
        if require_content and len(text) < 10:
            raise HTTPException(status_code=400, detail="题干过短（至少 10 字符）")
        fields["content"] = text
    for key in ("subject", "question_type", "knowledge_point", "answer",
                "analysis", "grade_band", "grade"):
        if key in payload:
            value = payload.get(key)
            fields[key] = (value or "").strip() if isinstance(value, str) else value
    if "knowledge_point" in fields:
        fields["knowledge_point"] = normalize_knowledge_point(
            fields["knowledge_point"], fields.get("subject", ""))
    if payload.get("difficulty") is not None:
        difficulty = int(payload["difficulty"])
        if not 1 <= difficulty <= 5:
            raise HTTPException(status_code=400, detail="difficulty 取值 1~5")
        fields["difficulty"] = difficulty
    if payload.get("options") is not None:
        options = payload["options"]
        if options and not isinstance(options, list):
            raise HTTPException(status_code=400, detail="options 应为 [{key,text}] 数组")
        fields["options"] = options
    if payload.get("knowledge_node_id") is not None:
        fields["knowledge_node_id"] = int(payload["knowledge_node_id"])
    if payload.get("status"):
        if payload["status"] not in ("draft", "pending", "published"):
            raise HTTPException(status_code=400, detail="status 取值 draft/pending/published")
        fields["status"] = payload["status"]
    return fields


def _validate_question_classification(fields: dict, existing: dict | None = None) -> None:
    """保证人工细分题型来自统一目录，避免页面与题库再次产生漂移。"""
    from edu_core.application.question_taxonomy import is_type_allowed

    subject = fields.get("subject", (existing or {}).get("subject", ""))
    question_type = fields.get("question_type", (existing or {}).get("question_type", ""))
    if not subject or not question_type:
        raise HTTPException(status_code=400, detail="subject 与 question_type 必填")
    if not is_type_allowed(subject, question_type):
        raise HTTPException(status_code=400, detail=f"{subject} 不支持题型：{question_type}")


@router.post("/teacher/questions")
def create_question(payload: dict, request: Request,
                    user: dict[str, Any] = Depends(require_teacher)) -> dict:
    """录入完整题目（AI 确认流保存入口）：题干/选项/答案/解析/难度/学段/知识点。"""
    fields = _full_fields(payload)
    _validate_question_classification(fields)
    stores = _stores()
    question_id = stores.questions.insert(
        fields["content"], fields.get("subject", ""),
        question_type=fields.get("question_type", ""),
        knowledge_point=fields.get("knowledge_point", ""),
        source="teacher", options=fields.get("options"),
        answer=fields.get("answer"), analysis=fields.get("analysis"),
        difficulty=fields.get("difficulty"), grade_band=fields.get("grade_band"),
        grade=fields.get("grade"), knowledge_node_id=fields.get("knowledge_node_id"),
        created_by=int(user["id"]), status=fields.get("status", "published"))
    stores.audit.insert(action="create_question", user_id=int(user["id"]), username=user.get("username"),
                        resource=f"teacher/questions/{question_id}", detail={"status": fields.get("status", "published")},
                        client_ip=request.client.host if request.client else None)
    return {"status": "ok", "id": question_id}


@router.post("/teacher/questions/batch")
def create_questions_batch(payload: dict, request: Request,
                           user: dict[str, Any] = Depends(require_teacher)) -> dict:
    """批量确认入库；先完整校验，再逐题保存并返回每题 ID。"""
    items = payload.get("items") or []
    if not items or len(items) > 50:
        raise HTTPException(status_code=400, detail="items 需为 1~50 条")
    fields_list = [_full_fields(item) for item in items]
    for fields in fields_list:
        _validate_question_classification(fields)

    stores = _stores()
    created: list[dict[str, int]] = []
    try:
        for index, fields in enumerate(fields_list):
            question_id = stores.questions.insert(
                fields["content"], fields.get("subject", ""),
                question_type=fields.get("question_type", ""),
                knowledge_point=fields.get("knowledge_point", ""), source="teacher",
                options=fields.get("options"), answer=fields.get("answer"), analysis=fields.get("analysis"),
                difficulty=fields.get("difficulty"), grade_band=fields.get("grade_band"),
                grade=fields.get("grade"), knowledge_node_id=fields.get("knowledge_node_id"),
                created_by=int(user["id"]), status=fields.get("status", "published"))
            created.append({"index": index, "id": question_id})
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"批量入库在第 {len(created) + 1} 题失败；已保存 {len(created)} 题：{str(exc)[:120]}") from exc
    stores.audit.insert(action="create_questions_batch", user_id=int(user["id"]), username=user.get("username"),
                        resource="teacher/questions/batch", detail={"count": len(created)},
                        client_ip=request.client.host if request.client else None)
    return {"status": "ok", "created": created, "total": len(created)}


@router.put("/teacher/questions/{question_id}")
def update_question(question_id: int, payload: dict, request: Request,
                    user: dict[str, Any] = Depends(require_teacher)) -> dict:
    """编辑题目（部分字段即可）。"""
    fields = _full_fields(payload, require_content=False)
    stores = _stores()
    existing = stores.questions.get(question_id)
    if not existing:
        raise HTTPException(status_code=404, detail="题目不存在")
    _validate_question_classification(fields, existing)
    stores.questions.update(question_id, **fields)
    stores.audit.insert(action="update_question", user_id=int(user["id"]), username=user.get("username"),
                        resource=f"teacher/questions/{question_id}", detail={"fields": sorted(fields)},
                        client_ip=request.client.host if request.client else None)
    return {"status": "ok"}


# ---------------------------------------------------------------------------
# AI 预标注（M1 确认流）：只推理不落库，教师确认后走 POST /teacher/questions
# ---------------------------------------------------------------------------

def _preannotate(text: str, service) -> dict:
    # 原文中的人工标注优先，不调用外部模型；分类模型只负责没有明确标注的字段。
    embedded = extract_embedded_metadata(text)
    clean_text = embedded["text"]
    raw = service.predictor.predict(clean_text)
    resolved = extract_embedded_metadata(text, subject=raw["subject"])
    metadata = resolved["metadata"]
    knowledge = normalize_knowledge_point(metadata.get("knowledge_point") or raw["knowledge_point"], raw["subject"])
    confidence = dict(raw["confidence"])
    if metadata.get("knowledge_point"):
        confidence["knowledge_point"] = 1.0
    return {
        "text": clean_text,
        "subject": raw["subject"],
        "question_type": raw["question_type"],
        "knowledge_point": knowledge,
        "grade_band": metadata.get("grade_band") or raw.get("grade_band"),
        "answer": metadata.get("answer", ""),
        "analysis": metadata.get("analysis", ""),
        "difficulty": metadata.get("difficulty") or suggest_difficulty(clean_text, raw["question_type"]),
        "metadata_sources": resolved["sources"],
        "confidences": confidence,
    }


@router.post("/teacher/questions/classify")
def classify_one(payload: dict,
                 _: dict[str, Any] = Depends(require_teacher)) -> dict:
    """单条 AI 预标注。"""
    text = (payload.get("text") or "").strip()
    if len(text) < 10:
        raise HTTPException(status_code=400, detail="题干过短（至少 10 字符）")
    return _preannotate(text, _stores_service())


@router.post("/teacher/questions/classify-batch")
def classify_batch(payload: dict,
                   _: dict[str, Any] = Depends(require_teacher)) -> dict:
    """批量 AI 预标注（前端智能切分后调用，上限 50 条）。"""
    texts = payload.get("texts") or []
    if not texts or len(texts) > 50:
        raise HTTPException(status_code=400, detail="texts 需为 1~50 条")
    service = _stores_service()
    results = [_preannotate((t or "").strip(), service) for t in texts if len((t or "").strip()) >= 10]
    return {"items": results}


def _stores_service():
    from edu_core.application.factory import get_classification_service
    return get_classification_service()


# ---------------------------------------------------------------------------
# 组卷与试卷库（M1）
# ---------------------------------------------------------------------------

@router.post("/teacher/papers/generate")
def generate_paper(payload: dict,
                   _: dict[str, Any] = Depends(require_teacher)) -> dict:
    """配比组卷：{subject, grade_band?, type_counts:{选择题:n,...}, knowledge?, difficulty?}。

    规则式从题库（published）按各题型配额随机抽题。
    """
    subject = (payload.get("subject") or "").strip()
    if not subject:
        raise HTTPException(status_code=400, detail="subject 必填")
    grade_band = (payload.get("grade_band") or "").strip()
    knowledge = (payload.get("knowledge") or "").strip()
    difficulty = payload.get("difficulty")
    type_counts = payload.get("type_counts") or {}
    try:
        normalized_counts = {str(key): int(value) for key, value in type_counts.items()}
    except (AttributeError, TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail="type_counts 必须是题型到数量的映射") from exc
    if (not normalized_counts or sum(normalized_counts.values()) < 1
            or any(value < 0 or value > 100 for value in normalized_counts.values())):
        raise HTTPException(status_code=400, detail="type_counts 至少配置一种题型的数量")
    type_counts = normalized_counts
    from edu_core.application.question_taxonomy import is_type_allowed
    invalid_types = [name for name in type_counts if not is_type_allowed(subject, name)]
    if invalid_types:
        raise HTTPException(
            status_code=400, detail=f"{subject} 不支持题型：{', '.join(invalid_types)}")
    stores = _stores()
    picked: list[dict] = []
    shortages: list[dict] = []
    import random

    rng = random.Random()
    for qtype, want in type_counts.items():
        want = int(want)
        if want < 1:
            continue
        items, _total = stores.questions.list(
            subject=subject, question_type=qtype, knowledge=knowledge,
            grade_band=grade_band, difficulty=difficulty,
            status="published", limit=500)
        if not items:
            shortages.append({"question_type": qtype, "requested": want, "available": 0})
            continue
        actual = min(want, len(items))
        picked.extend(rng.sample(items, actual))
        if actual < want:
            shortages.append({"question_type": qtype, "requested": want, "available": actual})
    if not picked:
        raise HTTPException(status_code=404, detail="题库中没有匹配的题目，请调整筛选条件")
    return {"questions": picked, "total": len(picked), "shortages": shortages}


@router.post("/teacher/papers")
def save_paper(payload: dict, request: Request, user: dict[str, Any] = Depends(require_teacher)) -> dict:
    """保存试卷：{title, subject, grade_band, question_ids:[...]}。"""
    title = (payload.get("title") or "").strip()
    subject = (payload.get("subject") or "").strip()
    grade_band = (payload.get("grade_band") or "高中").strip()
    question_ids = payload.get("question_ids") or []
    if len(set(int(q) for q in question_ids)) != len(question_ids):
        raise HTTPException(status_code=400, detail="question_ids 存在重复题目")
    if not title or not question_ids:
        raise HTTPException(status_code=400, detail="title 与 question_ids 必填")
    stores = _stores()
    paper_id = stores.papers.create(title, subject, grade_band,
                                    created_by=int(user["id"]),
                                    question_ids=[int(q) for q in question_ids])
    stores.audit.insert(action="create_paper", user_id=int(user["id"]), username=user.get("username"),
                        resource=f"teacher/papers/{paper_id}", detail={"question_count": len(question_ids)},
                        client_ip=request.client.host if request.client else None)
    return {"status": "ok", "id": paper_id}


@router.get("/teacher/papers")
def list_papers(limit: int = 50, offset: int = 0,
                user: dict[str, Any] = Depends(require_teacher)) -> dict:
    items, total = _stores().papers.list(created_by=int(user["id"]),
                                         limit=limit, offset=offset)
    return {"total": total, "items": items}


@router.get("/teacher/papers/{paper_id}")
def get_paper(paper_id: int, user: dict[str, Any] = Depends(require_teacher)) -> dict:
    paper = _stores().papers.get(paper_id)
    if not paper:
        raise HTTPException(status_code=404, detail="试卷不存在")
    if int(paper["created_by"]) != int(user["id"]):
        raise HTTPException(status_code=403, detail="只能查看自己创建的试卷")
    return paper


@router.delete("/teacher/papers/{paper_id}")
def delete_paper(paper_id: int, request: Request, user: dict[str, Any] = Depends(require_teacher)) -> dict:
    get_paper(paper_id, user)
    if not _stores().papers.delete(paper_id):
        raise HTTPException(status_code=404, detail="试卷不存在")
    _stores().audit.insert(action="delete_paper", user_id=int(user["id"]), username=user.get("username"),
                           resource=f"teacher/papers/{paper_id}", detail={},
                           client_ip=request.client.host if request.client else None)
    return {"status": "ok"}


@router.get("/teacher/papers/{paper_id}/export")
def export_paper(paper_id: int, request: Request, user: dict[str, Any] = Depends(require_teacher)) -> Response:
    """导出 Word 打印版：第一页试卷，第二页参考答案与解析。"""
    from docx import Document

    paper = get_paper(paper_id, user)
    doc = Document()
    doc.add_heading(paper["title"], level=0)
    doc.add_paragraph(f"学科：{paper['subject']}    学段：{paper['grade_band']}    "
                      f"共 {len(paper['questions'])} 题")
    for i, q in enumerate(paper["questions"], start=1):
        doc.add_paragraph(f"{i}.（{q['question_type']}）{q['content']}")
        if q.get("options"):
            for opt in q["options"]:
                doc.add_paragraph(f"    {opt.get('key')}. {opt.get('text')}")
    doc.add_page_break()
    doc.add_heading("参考答案与解析", level=1)
    for i, q in enumerate(paper["questions"], start=1):
        line = f"{i}. 答案：{q.get('answer') or '—'}"
        if q.get("knowledge_point"):
            line += f"    知识点：{q['knowledge_point']}"
        if q.get("difficulty"):
            line += f"    难度：{q['difficulty']}"
        doc.add_paragraph(line)
        if q.get("analysis"):
            doc.add_paragraph(f"    解析：{q['analysis']}")

    buf = BytesIO()
    doc.save(buf)
    from urllib.parse import quote

    filename = paper["title"].replace(" ", "_") + ".docx"
    _stores().audit.insert(action="export_paper", user_id=int(user["id"]), username=user.get("username"),
                           resource=f"teacher/papers/{paper_id}", detail={"format": "docx"},
                           client_ip=request.client.host if request.client else None)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition":
                 f"attachment; filename=\"paper_{paper_id}.docx\"; "
                 f"filename*=UTF-8''{quote(filename)}"})


# ---------------------------------------------------------------------------
# 作业发布与批改（M2 / C3）
# ---------------------------------------------------------------------------

@router.get("/teacher/assignments")
def list_assignments(limit: int = 50, offset: int = 0,
                     user: dict[str, Any] = Depends(require_teacher)) -> dict:
    from edu_core.application.factory import get_assignment_service

    return get_assignment_service().list_for_teacher(
        int(user["id"]), limit=limit, offset=offset)


@router.post("/teacher/assignments")
def create_assignment(payload: dict, request: Request,
                      user: dict[str, Any] = Depends(require_teacher)) -> dict:
    from edu_core.application.factory import get_assignment_service

    result = get_assignment_service().create_assignment(
        paper_id=payload.get("paper_id"), class_id=payload.get("class_id"),
        title=payload.get("title", ""), created_by=int(user["id"]),
        mode=payload.get("mode", "homework"), due_at=payload.get("due_at"),
        allow_self_check=payload.get("allow_self_check", True),
    )
    _stores().audit.insert(
        action="create_assignment", user_id=user.get("id") or None,
        username=user.get("username"), resource=f"assignments/{result['id']}",
        detail={"paper_id": result["paper_id"], "class_id": result["class_id"],
                "mode": result["mode"]},
        client_ip=request.client.host if request.client else None)
    return result


@router.get("/teacher/assignments/{assignment_id}")
def get_assignment(assignment_id: int,
                   user: dict[str, Any] = Depends(require_teacher)) -> dict:
    from edu_core.application.factory import get_assignment_service

    return get_assignment_service().get_for_teacher(assignment_id, int(user["id"]))


@router.get("/teacher/assignments/{assignment_id}/submissions")
def assignment_submissions(assignment_id: int,
                           user: dict[str, Any] = Depends(require_teacher)) -> dict:
    from edu_core.application.factory import get_assignment_service

    return get_assignment_service().list_submissions(assignment_id, int(user["id"]))


@router.post("/teacher/submissions/{submission_id}/check")
def check_submission(submission_id: int, payload: dict, request: Request,
                     user: dict[str, Any] = Depends(require_teacher)) -> dict:
    from edu_core.application.factory import get_assignment_service

    result = get_assignment_service().check_submission(
        submission_id, int(user["id"]), payload.get("final_score"))
    _stores().audit.insert(
        action="check_submission", user_id=user.get("id") or None,
        username=user.get("username"), resource=f"submissions/{submission_id}",
        detail={"final_score": result["final_score"]},
        client_ip=request.client.host if request.client else None)
    return result
