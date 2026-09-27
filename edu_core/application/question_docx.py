"""题目文件文本提取与服务端切分规则（DOCX / PDF / Markdown / TXT）。"""

from __future__ import annotations

import re
import zipfile
from io import BytesIO
from pathlib import Path

MAX_UNCOMPRESSED_BYTES = 80 * 1024 * 1024
STANDARD_START = re.compile(r"^\s*(?:第\s*)?(\d{1,3})\s*[.．、)]\s*")
BRACKET_START = re.compile(r"^\s*[（(]\s*(\d{1,3})\s*[)）]\s*")
TEXT_SUFFIXES = {".pdf", ".md", ".markdown", ".txt"}


def _safe_docx(payload: bytes) -> None:
    try:
        with zipfile.ZipFile(BytesIO(payload)) as archive:
            names = set(archive.namelist())
            total = sum(item.file_size for item in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise ValueError("文件不是有效的 DOCX 文档") from exc
    if "word/document.xml" not in names:
        raise ValueError("文件缺少 DOCX 正文结构")
    if total > MAX_UNCOMPRESSED_BYTES:
        raise ValueError("DOCX 解压后内容过大，拒绝处理")


def extract_docx_lines(payload: bytes) -> list[str]:
    _safe_docx(payload)
    try:
        from docx import Document

        document = Document(BytesIO(payload))
    except Exception as exc:
        raise ValueError("无法读取 DOCX 文档，请确认文件未损坏") from exc
    lines = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            value = "\t".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if value:
                lines.append(value)
    return lines


def split_questions(lines: list[str], *, max_questions: int = 50) -> list[str]:
    """优先按 ``1.``/``1、``切分；仅无此格式时才把 ``（1）`` 当题号。"""
    if not lines:
        return []
    pattern = STANDARD_START if any(STANDARD_START.match(line) for line in lines) else BRACKET_START
    blocks: list[list[str]] = []
    current: list[str] = []
    for line in lines:
        if pattern.match(line) and current:
            blocks.append(current)
            current = [line]
        else:
            current.append(line)
    if current:
        blocks.append(current)
    questions = ["\n".join(block).strip() for block in blocks]
    questions = [question for question in questions if len(question) >= 8]
    if len(questions) > max_questions:
        raise ValueError(f"文档识别出 {len(questions)} 题，单次最多 {max_questions} 题")
    return questions


def parse_docx_questions(payload: bytes, *, max_questions: int = 50) -> list[str]:
    return split_questions(extract_docx_lines(payload), max_questions=max_questions)


def parse_document_questions(filename: str, payload: bytes, *, max_questions: int = 50) -> list[str]:
    """按扩展名分派提取方式：DOCX 走段落/表格，PDF/MD/TXT 复用资料文本提取后切分。

    扫描型 PDF 在 OCR 未开启时会让资料提取明确报错，这里原样转成 400 语义。
    """
    suffix = Path(filename or "").suffix.lower()
    if suffix == ".docx":
        return parse_docx_questions(payload, max_questions=max_questions)
    if suffix in TEXT_SUFFIXES:
        from edu_core.rag.loaders import DocumentLoadError, load_document_bytes

        try:
            loaded = load_document_bytes(filename, payload)
        except DocumentLoadError as exc:
            raise ValueError(str(exc)) from exc
        lines = [line.strip() for line in loaded.text.splitlines() if line.strip()]
        return split_questions(lines, max_questions=max_questions)
    raise ValueError("仅支持 .docx / .pdf / .md / .txt 文件")
