"""教师题库 Excel 模板、导入校验、异步任务与导出。"""

from __future__ import annotations

import re
import threading
import zipfile
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from time import perf_counter
from uuid import uuid4

from edu_core.application.grading import grade_objective, is_objective_type
from edu_core.application.question_taxonomy import SUBJECTS, TYPE_META, is_type_allowed
from edu_core.storage.stores import QuestionStore

HEADERS = (
    "text", "subject", "question_type", "knowledge_point", "grade_band", "grade",
    "difficulty", "answer", "analysis", "options", "status",
)
REQUIRED_HEADERS = set(HEADERS[:-1])
STATUS_VALUES = {"published", "draft", "pending"}
GRADE_VALUES = {
    "初中": {"初一", "初二", "初三"},
    "高中": {"高一", "高二", "高三"},
}
MAX_ROWS = 5000
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024


def _openpyxl():
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - 部署依赖缺失时返回受控错误
        raise RuntimeError("Excel 功能依赖未安装，请执行 pip install -r requirements.txt") from exc
    return openpyxl


def _safe_xlsx(payload: bytes) -> None:
    """在解析前限制 ZIP 解压总量，避免小文件触发解压炸弹。"""
    try:
        with zipfile.ZipFile(BytesIO(payload)) as archive:
            total = sum(item.file_size for item in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise ValueError("文件不是有效的 xlsx 工作簿") from exc
    if total > MAX_UNCOMPRESSED_BYTES:
        raise ValueError("工作簿解压后内容过大，拒绝处理")


def _as_text(value) -> str:
    return "" if value is None else str(value).strip()


def _string_cell(cell, value: str) -> None:
    """强制以文本写出，防止题干以 =、+、-、@ 开头时成为 Excel 公式。"""
    cell.value = value
    cell.data_type = "s"


def _parse_options(value: str, *, allow_duplicate: bool = False) -> list[dict] | None:
    if not value:
        return None
    parts = [part.replace(r"\|", "|").strip()
             for part in re.split(r"(?<!\\)\|", value) if part.strip()]
    if not parts:
        return None
    options = []
    for index, part in enumerate(parts):
        match = re.match(r"^([A-H])\s*[.．、:]\s*(.+)$", part, re.IGNORECASE)
        key = match.group(1).upper() if match else chr(ord("A") + index)
        text = match.group(2).strip() if match else part
        if not allow_duplicate and any(item["key"] == key for item in options):
            raise ValueError(f"选项键 {key} 重复")
        options.append({"key": key, "text": text})
    expected = [chr(ord("A") + index) for index in range(len(options))]
    if not allow_duplicate and [item["key"] for item in options] != expected:
        raise ValueError("选项键必须从 A 开始连续排列")
    return options


def _format_options(options) -> str:
    if not options:
        return ""
    values = []
    for index, option in enumerate(options):
        if isinstance(option, dict):
            key = _as_text(option.get("key")) or chr(ord("A") + index)
            value = _as_text(option.get("text"))
            values.append(f"{key}. {value.replace('|', r'\|')}")
        else:  # 兼容早期导入的 ["A. ...", "B. ..."] 结构
            values.append(_as_text(option).replace("|", r"\|"))
    return "|".join(values)


def _style_sheet(sheet, *, rows: int, columns: int = 11, widths: list[int] | None = None) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    last_column = chr(64 + columns)
    sheet.auto_filter.ref = f"A1:{last_column}{max(2, rows)}"
    widths = widths or [48, 12, 18, 28, 10, 10, 10, 24, 48, 44, 12]
    for index, width in enumerate(widths[:columns], start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    header_fill = PatternFill("solid", fgColor="243B64")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=10, color="1F2937")
    line = Side(style="thin", color="DCE3EE")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=line)
    sheet.row_dimensions[1].height = 28
    for row in sheet.iter_rows(min_row=2, max_row=max(2, rows), min_col=1, max_col=columns):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=line)


def build_template() -> bytes:
    openpyxl = _openpyxl()
    from openpyxl.comments import Comment
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "题库导入"
    sheet.append(HEADERS)
    sheet.append(["", "", "", "", "", "", "", "", "", "", "published"])
    _style_sheet(sheet, rows=2)
    notes = {
        "A1": "必填。完整题干，至少 2 个字符。",
        "B1": "必填。数学、语文、英语、物理、化学、生物、历史、地理、政治。",
        "C1": "必填。题型须与学科匹配。",
        "D1": "知识点文本，如 一元二次方程。",
        "E1": "初中或高中。",
        "F1": "与学段匹配，如 初二、高三。",
        "G1": "可空；填写时为 1~5 的整数。",
        "H1": "选择题填 A/B 等；判断题填 正确/错误；主观题填参考答案。",
        "I1": "解题过程或说明，可留空。",
        "J1": "选择题使用 | 分隔，如 A. 甲|B. 乙|C. 丙|D. 丁。",
        "K1": "published、draft 或 pending；留空按 published。",
    }
    for address, note in notes.items():
        sheet[address].comment = Comment(note, "智慧教研平台")
    for address, values in {
        "B2:B5001": list(SUBJECTS),
        "C2:C5001": list(TYPE_META),
        "E2:E5001": ["初中", "高中"],
        "G2:G5001": ["1", "2", "3", "4", "5"],
        "K2:K5001": ["published", "draft", "pending"],
    }.items():
        validation = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
        validation.error = "请选择列表中的有效值"
        validation.errorTitle = "输入无效"
        validation.showErrorMessage = True
        sheet.add_data_validation(validation)
        validation.add(address)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_export(records: list[dict]) -> bytes:
    openpyxl = _openpyxl()
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "题库"
    sheet.append(HEADERS)
    for record in records:
        values = [
            _as_text(record.get("content")), _as_text(record.get("subject")),
            _as_text(record.get("question_type")), _as_text(record.get("knowledge_point")),
            _as_text(record.get("grade_band")), _as_text(record.get("grade")),
            record.get("difficulty"), _as_text(record.get("answer")),
            _as_text(record.get("analysis")), _format_options(record.get("options")),
            _as_text(record.get("status")),
        ]
        sheet.append(values)
        for column, value in enumerate(values, start=1):
            if isinstance(value, str):
                _string_cell(sheet.cell(sheet.max_row, column), value)
    _style_sheet(sheet, rows=sheet.max_row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_error_report(errors: list[dict]) -> bytes:
    openpyxl = _openpyxl()
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "错误行"
    sheet.append(["row", "field", "message"])
    for error in errors:
        values = [error.get("row"), _as_text(error.get("field")), _as_text(error.get("message"))]
        sheet.append(values)
        for column, value in enumerate(values, start=1):
            if isinstance(value, str):
                _string_cell(sheet.cell(sheet.max_row, column), value)
    _style_sheet(sheet, rows=sheet.max_row, columns=3, widths=[10, 22, 80])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def parse_workbook(payload: bytes) -> tuple[list[dict], list[dict]]:
    _safe_xlsx(payload)
    openpyxl = _openpyxl()
    try:
        workbook = openpyxl.load_workbook(BytesIO(payload), read_only=True, data_only=False)
    except Exception as exc:
        raise ValueError("无法读取工作簿，请确认文件未损坏且格式为 xlsx") from exc
    sheet = workbook["题库导入"] if "题库导入" in workbook.sheetnames else workbook.active
    raw_headers = [_as_text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    if len(raw_headers) != len(set(raw_headers)):
        raise ValueError("表头存在重复列")
    missing = sorted(REQUIRED_HEADERS - set(raw_headers))
    if missing:
        raise ValueError(f"缺少必要列：{', '.join(missing)}")
    positions = {name: raw_headers.index(name) for name in HEADERS if name in raw_headers}
    rows, errors = [], []
    seen_texts: set[str] = set()
    for row_number, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
        if row_number > MAX_ROWS + 1:
            errors.append({"row": row_number, "field": "row", "message": f"最多导入 {MAX_ROWS} 行"})
            break
        if not any(cell.value not in (None, "") for cell in cells):
            continue
        values = {}
        formula_field = ""
        for name, position in positions.items():
            cell = cells[position] if position < len(cells) else None
            if cell is not None and cell.data_type == "f":
                formula_field = name
                break
            values[name] = _as_text(cell.value if cell is not None else None)
        if formula_field:
            errors.append({"row": row_number, "field": formula_field, "message": "不允许使用公式单元格"})
            continue
        try:
            item = _validate_row(values)
            fingerprint = " ".join(item["content"].split())
            if fingerprint in seen_texts:
                raise ValueError("text: 与本文件前面的题干重复")
            seen_texts.add(fingerprint)
            rows.append(item)
        except ValueError as exc:
            message = str(exc)
            field = message.split(":", 1)[0] if ":" in message else "row"
            errors.append({"row": row_number, "field": field, "message": message})
    workbook.close()
    return rows, errors


def _validate_row(values: dict[str, str]) -> dict:
    content = values.get("text", "")
    subject = values.get("subject", "")
    question_type = values.get("question_type", "")
    if len(content) < 2:
        raise ValueError("text: 题干至少 2 个字符")
    if subject not in SUBJECTS:
        raise ValueError("subject: 不支持的学科")
    if not is_type_allowed(subject, question_type):
        raise ValueError("question_type: 题型与学科不匹配")
    grade_band = values.get("grade_band", "")
    grade = values.get("grade", "")
    if grade_band and grade_band not in GRADE_VALUES:
        raise ValueError("grade_band: 只能填写初中或高中")
    if grade and (not grade_band or grade not in GRADE_VALUES[grade_band]):
        raise ValueError("grade: 年级与学段不匹配")
    difficulty_text = values.get("difficulty", "")
    difficulty = None
    if difficulty_text:
        try:
            difficulty = int(difficulty_text)
        except ValueError as exc:
            raise ValueError("difficulty: 必须是 1~5 的整数") from exc
        if not 1 <= difficulty <= 5:
            raise ValueError("difficulty: 必须是 1~5 的整数")
    status = values.get("status") or "published"
    if status not in STATUS_VALUES:
        raise ValueError("status: 只能填写 published、draft 或 pending")
    try:
        options = _parse_options(
            values.get("options", ""),
            allow_duplicate=question_type not in {"选择题", "单选题", "多选题"},
        )
    except ValueError as exc:
        raise ValueError(f"options: {exc}") from exc
    if question_type in {"选择题", "单选题", "多选题"} and (not options or len(options) < 2):
        raise ValueError("options: 选择题至少填写两个选项")
    answer = values.get("answer", "")
    if is_objective_type(question_type):
        try:
            grade_objective(answer, answer, question_type)
        except ValueError as exc:
            raise ValueError(f"answer: {exc}") from exc
    return {
        "content": content, "subject": subject, "question_type": question_type,
        "knowledge_point": values.get("knowledge_point", "")[:128],
        "grade_band": grade_band or None, "grade": grade or None,
        "difficulty": difficulty, "answer": answer or None,
        "analysis": values.get("analysis", "") or None, "options": options, "status": status,
    }


class QuestionImportJobs:
    """进程内轻量异步任务；上传内容只驻留任务闭包，不落临时文件。"""

    def __init__(self, store: QuestionStore | None = None):
        self.store = store
        self._jobs: dict[str, dict] = {}
        self._lock = threading.Lock()
        self._executor = ThreadPoolExecutor(max_workers=2, thread_name_prefix="question-xlsx")

    def submit(self, payload: bytes, *, user_id: int, filename: str) -> str:
        job_id = uuid4().hex
        with self._lock:
            self._jobs[job_id] = {
                "job_id": job_id, "user_id": user_id, "filename": filename,
                "status": "QUEUED", "total": 0, "inserted": 0, "skipped": 0,
                "failed": 0, "errors": [], "elapsed_ms": None,
            }
        self._executor.submit(self._run, job_id, payload, user_id)
        return job_id

    def _run(self, job_id: str, payload: bytes, user_id: int) -> None:
        started = perf_counter()
        self._update(job_id, status="RUNNING")
        try:
            rows, errors = parse_workbook(payload)
            store = self.store or QuestionStore()
            inserted, skipped = store.bulk_insert(rows, created_by=user_id, source="excel_import")
            self._update(
                job_id, status="SUCCEEDED", total=len(rows) + len(errors),
                inserted=len(inserted), skipped=skipped, failed=len(errors), errors=errors,
                elapsed_ms=round((perf_counter() - started) * 1000),
            )
        except Exception as exc:  # noqa: BLE001 - 任务状态必须收敛且只暴露受控短消息
            self._update(job_id, status="FAILED", failed=1,
                         errors=[{"row": None, "field": "file", "message": str(exc)[:300]}],
                         elapsed_ms=round((perf_counter() - started) * 1000))

    def _update(self, job_id: str, **fields) -> None:
        with self._lock:
            self._jobs[job_id].update(fields)

    def get(self, job_id: str, user_id: int) -> dict | None:
        with self._lock:
            job = self._jobs.get(job_id)
            return dict(job) if job and int(job["user_id"]) == int(user_id) else None


_JOBS: QuestionImportJobs | None = None


def get_question_import_jobs() -> QuestionImportJobs:
    global _JOBS
    if _JOBS is None:
        _JOBS = QuestionImportJobs()
    return _JOBS
